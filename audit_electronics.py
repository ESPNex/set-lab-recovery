from pathlib import Path
import re,json,csv,hashlib,collections,shutil
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
R=Path(__file__).parent;O=R/'electronics-audit';O.mkdir(exist_ok=True)
# Documentary evidence read 2026-09-16. Only the stated subsystem is verified.
sources={
'Q_ELITE':{'url':'https://www.qualcomm.com/content/dam/qcomm-martech/dm-assets/documents/Snapdragon-8-Elite-Platform-Product-Brief.pdf','citationId':3,'publisher':'Qualcomm','verified':'SM8750-AB; LPDDR5x; up to 24 GB; UFS 4.0.','scope':'Product brief, not pinout/PCB/BOM'},
'Q_GEN3':{'url':'https://docs.qualcomm.com/doc/87-71408-1/87-71408-1_REV_G_Snapdragon_8_gen_3_Mobile_Platform_Product_Brief.pdf','citationId':1,'publisher':'Qualcomm','verified':'LPDDR5x up to 4800 MHz; up to 24 GB; UFS 4.0.','scope':'Product brief'},
'Q_7PLUS':{'url':'https://docs.qualcomm.com/doc/87-73943-1/87-73943-1_REV_E_Snapdragon_7__Gen_3_Mobile_Platform_Product_Brief.pdf','citationId':1,'publisher':'Qualcomm','verified':'LPDDR5x up to 4200 MHz; up to 24 GB; UFS 4.0.','scope':'Product brief'},
'Q_8S':{'url':'https://www.qualcomm.com/content/dam/qcomm-martech/dm-assets/documents/Product-Brief-Snapdragon-8s-Gen-4.pdf','citationId':1,'publisher':'Qualcomm','verified':'SM8735; LPDDR5x up to 4800 MHz; up to 24 GB; UFS 4.0.','scope':'Product brief'},
'M_G99':{'url':'https://www.mediatek.com/products/smartphones/mediatek-helio-g99','citationId':1,'publisher':'MediaTek','verified':'LPDDR4X; 4266 Mbps; UFS 2.2.','scope':'Official product specifications'},
'M_8300':{'url':'https://www.mediatek.com/products/tablets/mediatek-dimensity-8300','citationId':1,'publisher':'MediaTek','verified':'LPDDR5X up to 8533 Mbps; UFS 4.0 + MCQ; Cortex-A715/A510; Mali-G615.','scope':'8300 family; does not validate Ultra-specific PCB/package'},
'M_9400':{'url':'https://www.mediatek.com/products/smartphones/mediatek-dimensity-9400','citationId':3,'publisher':'MediaTek','verified':'LPDDR5X 10667 Mbps; UFS 4 + MCQ.','scope':'Official product specifications'},
'M_9500':{'url':'https://www.mediatek.com/products/smartphones/mediatek-dimensity-9500','citationId':1,'publisher':'MediaTek','verified':'TSMC N3P; Arm C1-Ultra/Premium/Pro; Mali-G1 Ultra MC12; LPDDR5X 10667; UFS 4.1 4-lane.','scope':'Official page, fetched chunks 1 and 2'},
'S_2500':{'url':'https://semiconductor.samsung.com/processor/mobile-processor/exynos-2500/','citationId':1,'publisher':'Samsung Semiconductor','verified':'LPDDR5X; UFS 4.0; 3nm GAA.','scope':'Official page, fetched specifications in chunk 3'},
'H_MEMORY':{'url':'https://product.skhynix.com/products/dram/lpddr.go?appTypCd=APX02&treeNo=1055','citationId':2,'publisher':'SK hynix','verified':'LPDDR5T/5X speed 8533–9600 Mbps; PoP/MCP/Discrete package types.','scope':'Memory family, not exact 16 GB MPN'},
'T_BGA':{'url':'https://www.ti.com/lit/an/spraa99c/spraa99c.pdf','citationId':2,'publisher':'Texas Instruments','verified':'Package-specific PCB land patterns; sub-mm pitch examples; electrical checks and X-ray inspection discussed.','scope':'Generic BGA process reference, not smartphone footprint specification'},
'U_PD':{'url':'https://www.usb.org/sites/default/files/2021-05/USB%20PG%20USB%20PD%203.1%20DevUpdate%20Announcement_FINAL.pdf','citationId':1,'publisher':'USB Promoter Group','verified':'PD3.1 extends power to 240 W and adds 28/36/48 V fixed voltages.','scope':'Protocol capability, not certification of catalog connector/charger'}
}
rules={'soc-x90':('LPDDR5X','Q_ELITE'),'soc-8e':('LPDDR5X','Q_GEN3'),'soc-7g2':('LPDDR5X','Q_7PLUS'),'soc-8s':('LPDDR5X','Q_8S'),'soc-g99':('LPDDR4X','M_G99'),'soc-d8300':('LPDDR5X','M_8300'),'soc-d9400':('LPDDR5X','M_9400'),'soc-dim95':('LPDDR5X','M_9500'),'soc-ex2500':('LPDDR5X','S_2500')}
s=(R/'uploads/ls.ts').read_text();q=r'(?:"(?:\\.|[^"\\])*"|\'(?:\\.|[^\'\\])*\')';pat=re.compile(r'^\s*p\(('+q+r'),\s*('+q+r'),\s*('+q+r'),\s*('+q+r'),\s*('+q+r')',re.M)
def unquote(x):return x[1:-1].replace('\\"','"').replace("\\'","'")
coords=json.loads((R/'coordinates-v17.json').read_text());ci=coords['items'];rows=[]
needed={'motherboard':'MPN/revisione PCB, schema, netlist, stack-up, Gerber/ODB++, BOM, XY pick-and-place, pinout, firmware e collaudi.','soc':'Package/revisione e ball map, reference design, PMIC e sequenze, BSP/firmware, lista memorie qualificate.','ram':'MPN completo, package PoP/MCP/discreto, densita die, canali, tensioni, timing, bin velocita e qualifica sul SoC.','storage':'MPN, ball map, tensioni, modalita host/device, boot e firmware; UFS/eMMC non deducibili dal solo rettangolo.','battery':'MPN e quote, topologia celle, protezione/BMS, NTC, Vmax di carica, correnti ammesse e caricatore qualificato.','display':'MPN pannello+touch, disegno quotato, connettore, pinout, MIPI DSI, rail, sequenze e driver.','cameraRear':'MPN modulo completo, ottica/OIS, connettore e CSI, rail, clock, driver e tarature ISP.','cameraFront':'MPN modulo completo, ottica/AF, connettore e CSI, rail, clock, driver e tarature ISP.','usb':'MPN connettore/flex, pinout, rating elettrico, controller PD, mux/retimer, ESD e progetto del power path.','antenna':'MPN antenna e radio, matching RF, bande, geometria massa, connettori e risultati RF.','speaker':'MPN, impedenza, potenza e risposta, codec/amplificatore e volume acustico.','mic':'MPN, interfaccia, bias/rail, clock, disposizione e DSP.','haptics':'MPN, risonanza/tensione, driver e meccanica.','sim':'MPN e quote del socket/carrello; pinout e controller.','thermal':'Disegno/spessore, compressione, conducibilita, contatti e verifica termica.','frame':'Disegno CAD e tolleranze con un modello/revisione identificato.','backcover':'Disegno CAD, aperture camere, materiali RF, tolleranze, fissaggi.','glass':'MPN vetro sagomato, quote, foro, adesivo, tolleranze e pannello abbinato.'}
for m in pat.finditer(s):
 sku,name,brand,cat,slot=[unquote(m.group(j)) for j in range(1,6)];src='';status='DATI_INSUFFICIENTI';note='Voce di catalogo non equivalente a part number acquistabile e qualificato. Non ricercata individualmente in questa prima passata.'
 if sku in rules:src=rules[sku][1];status='FAMIGLIA_DOCUMENTATA_NON_QUALIFICATA';note='Verificata famiglia e tipo memoria pubblicato; non package/pinout, motherboard o tutti i dati del catalogo.'
 if sku=='soc-q6':src='M_8300';note='Ricercata famiglia 8300; suffisso Ultra e implementazione specifica non qualificati.'
 if slot=='motherboard':status='PCB_NON_DOCUMENTATO';note='Nessun progetto PCB/fabbricazione individuato nei file del repository. Le ricerche mirate SBS-A7 e SBS-X9/SBS-G8 non hanno restituito documentazione pertinente; non prova di inesistenza.'
 if slot=='ram':src='H_MEMORY';note='Specifica generica senza MPN. Le cifre del catalogo in MHz richiedono revisione di clock vs data rate; la fonte SK hynix e di famiglia, non prova di ogni SKU.'
 if sku=='soc-dim95':note+=' Contrasto catalogo: 2 nm/Cortex-X930/Immortalis G930 contro N3P/C1/Mali-G1 Ultra nella fonte ufficiale.'
 if sku=='ram-16':note+=' La correzione grafica V17 9600 MHz era editoriale, non tecnicamente validata; 9600 Mbps e la misura di data rate pubblicata per la famiglia.'
 if sku.startswith('cons-'):status='UTENSILE_O_CONSUMABILE';note='Non componente elettronico saldato; richiede eventuale scheda tecnica/SDS e idoneita di processo.'
 if sku in ['disp-fold','glass-utg']:status='ESCLUSO_DALL_UTENTE';note='Escluso dalla selezione attiva non-fold.'
 sets=sorted({i['set'] for i in ci if i['sku']==sku})
 rows.append({'sku':sku,'nomeCatalogo':name,'marcaCatalogo':brand,'slot':slot,'set':','.join(map(str,sets)),'esitoDocumentale':status,'compatibilitaScheda':'NON_VERIFICATA' if not sku.startswith('cons-') else 'NON_APPLICABILE','collaudoFisico':'NON_ESEGUITO','nota':note,'datiNecessari':needed.get(slot,'MPN e documenti specifici'),'fonti':' ; '.join(sources[k]['url'] for k in src.split(';') if k),'verificaWebIndividuale': 'PARZIALE_FAMIGLIA' if sku in rules else 'NO_O_SOLO_FAMIGLIA'})
assert len(rows)==149;by={x['sku']:x for x in rows};filemap={i['file']:i['sku'] for i in ci};man=json.loads((R/'finals/composition-manifest.json').read_text());configs=[]
for c in man['compositions']:
 selected={role:filemap[Path(p).name] for role,p in c['selected'].items()};soc,ram,sto=selected['soc'],selected['ram'],selected['storage'];rt=re.search(r'LPDDR\dX?',by[ram]['nomeCatalogo']);actual=rt.group(0) if rt else 'UNKNOWN';status='NON_VALUTATA_SUFFICIENTEMENTE';finding='Nessuna conclusione sulla compatibilita del core: ricerca specifica ancora incompleta.';url=''
 if soc in rules:
  req,source=rules[soc];url=sources[source]['url']
  if actual!=req:status='CONFLITTO_RAM_DOCUMENTATO';finding=f'{soc}: memoria pubblicata {req}; selezionata {ram} {actual}. Non approvabile come abbinamento del controller documentato.'
  else:status='SOLO_TIPO_RAM_COINCIDENTE';finding='Tipo LPDDR coincide; MPN/capacita/bin/timing/package, storage, PCB e sistema non qualificati.'
 if soc=='soc-g99':finding+=' Host documentato UFS 2.2: nessuna modalita UFS 4.0 garantita per sto-1tb; retrocompatibilita del dispositivo esatto da verificare, non assumere incompatibilita assoluta dalla sola versione.'
 if soc=='soc-q6':finding+=' Famiglia 8300 LPDDR5X: chiarire Ultra, velocita operativa e MPN; non estesa automaticamente la qualifica della famiglia.';url=sources['M_8300']['url']
 configs.append({'immagine':c['file'],'coppia':'-'.join(f'{s:02d}' for s in c['pair']),'soc':soc,'ram':ram,'storage':sto,'esitoCore':status,'esitoConfigurazione':'BLOCCATA_PER_PROGETTAZIONE_REALE','motivo':finding,'bloccoComune':'Mainboard senza schema/netlist/footprint/BOM/XY verificati; nessun collaudo reale.','fonte':url})
 c['electronicAudit']={'status':'NOT_VALIDATED_DO_NOT_BUILD','coreFinding':status,'detail':finding,'source':url,'auditReport':'../electronics-audit/AUDIT-ELETTRONICO.md'}
summary={'date':'2026-09-16','reviewedAssetCommit':'7628719efdb77f3bf52b30f8282554a1e3bc2a6f','catalogSku':149,'activeSku':147,'documentedSocFamilies':len(rules),'configurations':29,'documentedRamConflicts':sum(x['esitoCore']=='CONFLITTO_RAM_DOCUMENTATO' for x in configs),'validatedConfigurations':0,'physicalTestsPerformed':0,'newAssemblyImagesCreated':0,'scope':'Preliminary documentary screening; all SKU catalog rows inventoried, only selected families researched online. No complete 149-part datasheet qualification.'}
(O/'sources.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2));(O/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2));(R/'finals/composition-manifest.json').write_text(json.dumps(man,ensure_ascii=False,indent=2))
for name,data in [('sku-audit.csv',rows),('configurations-audit.csv',configs)]:
 with (O/name).open('w') as out:w=csv.DictWriter(out,fieldnames=data[0].keys(),lineterminator='\n');w.writeheader();w.writerows(data)
wb=Workbook();wb.remove(wb.active)
for name,data in [('149 SKU',rows),('29 configurazioni',configs),('Fonti',[{'id':k,**v} for k,v in sources.items()]),('Riepilogo',[{'metrica':k,'valore':str(v)} for k,v in summary.items()])]:
 ws=wb.create_sheet(name);ws.append(list(data[0].keys()))
 for row in data:ws.append(list(row.values()))
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='17392E')
 for col in ws.columns:
  k=col[0].column_letter;ws.column_dimensions[k].width=min(65,max(18,len(str(col[0].value))+4))
 for row in ws.iter_rows(min_row=2):
  for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
  ws.row_dimensions[row[0].row].height=65
wb.save(O/'AUDIT-ELETTRONICO.xlsx');print(summary)
